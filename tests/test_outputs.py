"""End-to-end smoke test: a small run must produce a consistent workbook, memo and README block."""

import numpy as np
import pytest
from openpyxl import load_workbook

from captive import excel, figures, memo
from captive.analysis import Study
from captive.cli import summary_markdown, update_readme
from captive.config import with_overrides


@pytest.fixture(scope="module")
def small_run(base_cfg, tmp_path_factory):
    cfg = with_overrides(base_cfg, {"meta.n_sims": 3000, "breakeven.sims": 2000, "breakeven.scales": [1.0, 2.0, 4.0]})
    cfg["sensitivity"] = cfg["sensitivity"][:2]
    out = Study(cfg).run()
    tmp = tmp_path_factory.mktemp("run")
    figs = figures.make_all(out, tmp / "figures")
    return cfg, out, figs, tmp


def test_expected_cashflows_discount_to_the_mean_pv(small_run):
    """The Excel NPV formula discounts expected flows; by linearity it must equal the mean path PV."""
    cfg, out, _, _ = small_run
    r = cfg["economics"]["discount_rate"]
    for p in out["programs"]:
        flows = np.array(p["expected_cashflows"])
        npv = flows @ (1 + r) ** -np.arange(len(flows))
        assert npv == pytest.approx(p["mean"], rel=1e-9)


def test_recommendation_is_consistent_with_numbers(small_run):
    _, out, _, _ = small_run
    rec = out["recommendation"]
    best = min(out["programs"], key=lambda p: p["mean"])
    assert rec["best_program"] == best["id"]
    assert rec["decision"] in {"proceed", "large_deductible_now_revisit_captive", "stay_guaranteed_cost"}


def test_workbook_has_all_sheets_and_live_formulas(small_run):
    cfg, out, figs, tmp = small_run
    path = excel.build_workbook(out, cfg, tmp / "model.xlsx", figs)
    wb = load_workbook(path)
    for name in ["Cover", "Results", "Cash Flows", "Assumptions", "Loss Model", "Capital", "Domiciles",
                 "Sensitivity", "Scenarios", "Break-even", "Sim Sample", "Figures"]:
        assert name in wb.sheetnames
    cf = wb["Cash Flows"]
    formulas = [c.value for row in cf.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("SUMPRODUCT" in f for f in formulas)
    assert any("IRR(" in f for f in formulas)
    statuses = {c.value for c in wb["Assumptions"]["C"] if c.value}
    assert {"verified", "assumption"} <= statuses


def test_memo_renders_markdown_and_pdf(small_run):
    cfg, out, _, tmp = small_run
    blocks = memo.build_blocks(out, cfg, tmp / "figures")
    md = memo.to_markdown(blocks)
    assert out["recommendation"]["headline"] in md
    assert "hypothetical" in md.lower()
    pdf = memo.to_pdf(blocks, tmp / "memo.pdf")
    assert pdf.stat().st_size > 20_000
    assert pdf.read_bytes()[:4] == b"%PDF"


def test_readme_block_is_replaced_between_markers(small_run, tmp_path):
    cfg, out, _, _ = small_run
    readme = tmp_path / "README.md"
    readme.write_text("intro\n<!-- RESULTS:START -->\nold\n<!-- RESULTS:END -->\noutro\n", encoding="utf-8")
    assert update_readme(summary_markdown(out, cfg), readme)
    text = readme.read_text(encoding="utf-8")
    assert "old" not in text and "Model decision" in text and text.endswith("outro\n")
