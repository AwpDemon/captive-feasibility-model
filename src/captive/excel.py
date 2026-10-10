"""Build the broker-style Excel workbook from a results dictionary (openpyxl).

The workbook is generated, never hand-edited, so it always matches the Python run.
The Cash Flows sheet keeps live formulas (discount factors, NPV, IRR) driven by an
editable discount-rate cell, so an Excel user can re-discount the expected flows.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import flatten

NAVY = "1F3A5F"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
NOTE_FONT = Font(italic=True, size=9, color="52514E")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
STATUS_FILL = {
    "verified": PatternFill("solid", fgColor="D9EAD3"),
    "benchmark": PatternFill("solid", fgColor="DDEBF7"),
    "secondary": PatternFill("solid", fgColor="FCE4D6"),
    "assumption": PatternFill("solid", fgColor="EDEDED"),
}
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
USD = '"$"#,##0;[Red]-"$"#,##0'
PCT = "0.0%"


def _title(ws: Worksheet, text: str, sub: str | None = None) -> int:
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    row = 2
    if sub:
        ws["A2"] = sub
        ws["A2"].font = NOTE_FONT
        row = 3
    return row + 1


def _table(ws: Worksheet, top: int, headers: list[str], rows: list[list], formats: dict[int, str] | None = None,
           left: int = 1) -> tuple[int, int]:
    """Write a header + rows block; returns (first data row, last data row)."""
    formats = formats or {}
    for j, h in enumerate(headers):
        c = ws.cell(row=top, column=left + j, value=h)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r):
            if isinstance(v, (np.floating, np.integer)):
                v = v.item()
            c = ws.cell(row=top + i, column=left + j, value=v)
            c.border = BOX
            if j in formats:
                c.number_format = formats[j]
    return top + 1, top + len(rows)


def _widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build_workbook(out: dict, cfg: dict, path: Path, figures: dict[str, Path] | None = None) -> Path:
    wb = Workbook()
    progs = out["programs"]
    pid_label = {p["id"]: p["label"] for p in progs}
    rec = out["recommendation"]
    meta = out["meta"]

    # ------------------------------------------------------------------ Cover
    ws = wb.active
    ws.title = "Cover"
    r = _title(ws, "Captive Feasibility Model - HypoMfg, Inc. (HYPOTHETICAL)",
               "Illustrative study for a hypothetical company. Not an actuarial opinion, tax or legal advice.")
    lines = [
        ("Recommendation", rec["headline"]),
        ("Lowest expected-cost program", rec["best_label"]),
        ("Company", f"{cfg['company']['description']}; revenue ${cfg['company']['revenue']:,.0f}; "
                    f"{cfg['company']['employees']} employees"),
        ("Monte Carlo", f"{meta['n_sims']:,} paths x {meta['horizon_years']} policy years + run-off; seed {meta['seed']}"),
        ("Discount rate", f"{meta['discount_rate']:.1%} (after-tax opportunity cost of cash)"),
        ("Model version", meta["model_version"]),
        ("Generated", date.today().isoformat()),
        ("Reproduce", "pip install -e . ; captive run   (all inputs: config/assumptions.yaml)"),
    ]
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Sheets").font = Font(bold=True, color=NAVY)
    toc = [
        ("Results", "Program comparison: PV of 5-year after-tax cost, percentiles, probability of beating GC/LD"),
        ("Cash Flows", "Expected cash flows by year with LIVE NPV/IRR formulas (edit the yellow discount-rate cell)"),
        ("Assumptions", "Every input from the YAML config with provenance status and source"),
        ("Loss Model", "Frequency/severity by line; analytic vs simulated expected losses"),
        ("Capital", "Retained-loss percentiles (1-in-100, 1-in-200), surplus paths, capital calls, stress tests"),
        ("Domiciles", "Vermont, North Carolina, Utah, Bermuda, Cayman: regulatory figures with citations + cost"),
        ("Sensitivity", "Tornado: one driver at a time, low/high"),
        ("Scenarios", "Whole-world alternatives (hard market, adverse losses, tax treatment, size)"),
        ("Break-even", "Saving vs guaranteed cost as the company grows"),
        ("Sim Sample", "First 2,000 simulated paths for your own pivots"),
        ("Figures", "Charts from the Python run"),
    ]
    for name, desc in toc:
        r += 1
        c = ws.cell(row=r, column=1, value=name)
        c.hyperlink = f"#'{name}'!A1"
        c.font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=2, value=desc)
    _widths(ws, {"A": 30, "B": 110})

    # ------------------------------------------------------------------ Results
    ws = wb.create_sheet("Results")
    r = _title(ws, "Program comparison (present value of 5-year after-tax cost of risk, lower is better)",
               "Each program evaluated on the same simulated claims. Savings are mean differences across paths.")
    headers = ["Program", "Type", "Mean PV cost", "Std dev", "5th pct", "25th pct", "Median", "75th pct", "95th pct",
               "99th pct", "Saving vs GC", "P(cheaper than GC)", "Saving vs LD", "P(cheaper than LD)",
               "Year-1 premium outlay"]
    rows = [[p["label"], p["type"], p["mean"], p["std"], p["p5"], p["p25"], p["p50"], p["p75"], p["p95"], p["p99"],
             p["saving_vs_gc_mean"], p["prob_cheaper_than_gc"], p.get("saving_vs_ld_mean"),
             p.get("prob_cheaper_than_ld"), p["year1_premium_outlay"]] for p in progs]
    fmt = {i: USD for i in range(2, 11)} | {11: PCT, 12: USD, 13: PCT, 14: USD}
    first, last = _table(ws, r, headers, rows, fmt)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Mean PV cost by program"
    chart.y_axis.title = "PV of 5-year after-tax cost ($)"
    chart.y_axis.numFmt = '"$"#,##0'
    chart.add_data(Reference(ws, min_col=3, min_row=r, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.legend = None
    chart.height, chart.width = 8, 18
    ws.add_chart(chart, f"A{last + 3}")

    r2 = last + 22
    ws.cell(row=r2, column=1, value="Mean PV by cost component (sums to the mean PV cost)").font = Font(bold=True, color=NAVY)
    comps = list(progs[0]["components"].keys())
    _table(ws, r2 + 1, ["Component"] + [p["label"] for p in progs],
           [[c] + [p["components"][c] for p in progs] for c in comps]
           + [["Total"] + [p["mean"] for p in progs]], {j: USD for j in range(1, len(progs) + 1)})
    _widths(ws, {get_column_letter(i): 15 for i in range(2, 16)} | {"A": 42})
    ws.freeze_panes = ws.cell(row=r + 1, column=2)

    # --------------------------------------------------------------- Cash Flows
    ws = wb.create_sheet("Cash Flows")
    r = _title(ws, "Expected parent cash flows by period (cost positive; income-tax effects included)",
               "Period 0 = first policy inception. Losses run off to full payment. Flows are path averages from the "
               "simulation; NPV/IRR below are live Excel formulas.")
    ws.cell(row=r, column=1, value="Discount rate (edit me)").font = Font(bold=True)
    rate_cell = ws.cell(row=r, column=2, value=meta["discount_rate"])
    rate_cell.number_format, rate_cell.fill = PCT, INPUT_FILL
    r += 2
    n_t = len(progs[0]["expected_cashflows"])
    header_row = r
    _table(ws, r, ["Period"] + list(range(n_t)) + ["NPV (formula)", "Python PV (check)"], [])
    r += 1
    df_row = r
    ws.cell(row=r, column=1, value="Discount factor")
    for t in range(n_t):
        col = get_column_letter(2 + t)
        ws.cell(row=r, column=2 + t, value=f"=1/(1+$B${rate_cell.row})^{col}${header_row}").number_format = "0.0000"
    r += 1
    first_flow = r
    for p in progs:
        ws.cell(row=r, column=1, value=p["label"])
        for t, v in enumerate(p["expected_cashflows"]):
            ws.cell(row=r, column=2 + t, value=float(v)).number_format = USD
        lc = get_column_letter(1 + n_t)
        ws.cell(row=r, column=2 + n_t, value=f"=SUMPRODUCT(B{r}:{lc}{r},B${df_row}:{lc}${df_row})").number_format = USD
        ws.cell(row=r, column=3 + n_t, value=p["mean"]).number_format = USD
        r += 1
    r += 1
    focus = meta["focus_captive"]
    ws.cell(row=r, column=1, value=f"Incremental cash saved by {pid_label[focus]} vs guaranteed cost").font = Font(bold=True)
    r += 1
    gc_row = first_flow + [p["id"] for p in progs].index("gc")
    f_row = first_flow + [p["id"] for p in progs].index(focus)
    ws.cell(row=r, column=1, value="GC flow minus captive flow")
    for t in range(n_t):
        col = get_column_letter(2 + t)
        ws.cell(row=r, column=2 + t, value=f"={col}{gc_row}-{col}{f_row}").number_format = USD
    lc = get_column_letter(1 + n_t)
    ws.cell(row=r + 1, column=1, value="NPV of incremental flows")
    ws.cell(row=r + 1, column=2, value=f"=SUMPRODUCT(B{r}:{lc}{r},B${df_row}:{lc}${df_row})").number_format = USD
    ws.cell(row=r + 2, column=1, value="IRR of incremental flows (Excel)")
    ws.cell(row=r + 2, column=2, value=f'=IFERROR(IRR(B{r}:{lc}{r}),"n/a - no sign change")').number_format = PCT
    py_irr = out["incremental"]["vs_gc_irr"]
    ws.cell(row=r + 3, column=1, value="IRR (Python check)")
    ws.cell(row=r + 3, column=2, value=py_irr if py_irr is not None else "n/a").number_format = PCT
    ws.cell(row=r + 5, column=1, value="Note: NPV of expected flows equals the mean of path NPVs because NPV is linear.").font = NOTE_FONT
    _widths(ws, {"A": 44} | {get_column_letter(i): 12 for i in range(2, n_t + 4)})

    # -------------------------------------------------------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    r = _title(ws, "All model inputs (generated from config/assumptions.yaml)",
               "Status: verified = statute/regulator/official data; benchmark = public figure used to anchor an "
               "assumption; secondary = vendor/association; assumption = judgment. Retrieved "
               f"{cfg.get('sources', {}).get('_retrieved', 'n/a')}.")
    sources = cfg.get("sources", {})
    skip = ("sources", "sensitivity", "scenarios", "stress_tests", "breakeven")
    rows = []
    for path_, val in flatten({k: v for k, v in cfg.items() if k not in skip}):
        src = None
        for key in sorted(sources, key=len, reverse=True):
            if key != "_retrieved" and (path_ == key or path_.startswith(key + ".")):
                src = sources[key]
                break
        status = src["status"] if src else "assumption"
        rows.append([path_, val if not isinstance(val, (list, dict)) else str(val), status,
                     src.get("value_basis", "") if src else "", src.get("url", "") if src else ""])
    first, last = _table(ws, r, ["Input (YAML path)", "Value", "Status", "Basis / quote", "Source URL"], rows)
    for i in range(first, last + 1):
        st = ws.cell(row=i, column=3).value
        ws.cell(row=i, column=3).fill = STATUS_FILL.get(st, STATUS_FILL["assumption"])
        ws.cell(row=i, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        url = ws.cell(row=i, column=5).value
        if url:
            ws.cell(row=i, column=5).hyperlink = url
            ws.cell(row=i, column=5).font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A{r}:E{last}"
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    _widths(ws, {"A": 58, "B": 22, "C": 12, "D": 90, "E": 60})

    # ---------------------------------------------------------------- Loss Model
    ws = wb.create_sheet("Loss Model")
    r = _title(ws, "Loss model by line (policy year 1, ground-up within the per-occurrence limit)",
               "Analytic = closed-form limited expected value x expected count; simulated = Monte Carlo mean. "
               "The two agreeing is a built-in check of the simulator.")
    L = out["losses"]
    rows = [[x["label"], x["expected_claims_yr1"], x["simulated_claims_yr1"], x["expected_loss_yr1_analytic"],
             x["simulated_loss_yr1_mean"], x["simulated_loss_yr1_std"], x["simulated_loss_yr1_p99"], x["gc_premium_yr1"],
             x["limit"]] for x in L["lines"]]
    first, last = _table(ws, r, ["Line", "Expected claims", "Simulated claims", "Expected loss (analytic)",
                                 "Simulated mean loss", "Std dev", "99th pct", "GC premium (yr 1)", "Per-occ limit"],
                         rows, {1: "0.00", 2: "0.00", 3: USD, 4: USD, 5: USD, 6: USD, 7: USD, 8: USD})
    r = last + 2
    ws.cell(row=r, column=1, value="Components").font = Font(bold=True, color=NAVY)
    comp_rows = []
    for x in L["lines"]:
        for comp in x["components"]:
            f, s = comp["frequency"], comp["severity"]
            sev = (f"lognormal(mean ${s['mean']:,.0f}, CV {s['cv']})" if s["dist"] == "lognormal"
                   else f"Pareto(xm ${s['xm']:,.0f}, alpha {s['alpha']})")
            freq = f"{f.get('dist', 'negbin')}(mean {f['mean']}" + (f", VMR {f['vmr']})" if "vmr" in f else ")")
            comp_rows.append([x["label"], comp["name"], freq, sev])
    _table(ws, r + 1, ["Line", "Component", "Frequency", "Severity"], comp_rows)
    t = L["total_yr1"]
    r = r + len(comp_rows) + 3
    ws.cell(row=r, column=1, value="All lines combined, year 1").font = Font(bold=True, color=NAVY)
    _table(ws, r + 1, ["Mean", "Std dev", "Median", "90th", "95th", "99th", "99.5th"],
           [[t["mean"], t["std"], t["p50"], t["p90"], t["p95"], t["p99"], t["p99.5"]]], {i: USD for i in range(7)})
    _widths(ws, {"A": 30, "B": 18, "C": 34, "D": 40, "E": 18, "F": 14, "G": 14, "H": 16, "I": 14})

    # ------------------------------------------------------------------- Capital
    ws = wb.create_sheet("Capital")
    cap = out["capital"]
    r = _title(ws, f"Capital adequacy - {pid_label[focus]} ({cfg['domiciles'][cfg['captive']['base_domicile']]['name']})",
               "Capital = max(domicile minimum, 1-in-100 year-1 retained loss beyond the funded loss pick, premium/3).")
    a = cap["annual_retained"]
    kv = [
        ("Capital contributed", cap["capital"], USD),
        ("  basis: domicile minimum", cap["capital_basis"]["domicile_minimum"], USD),
        ("  basis: VaR-based", cap["capital_basis"]["var_based"], USD),
        ("  basis: premium-to-surplus", cap["capital_basis"]["premium_to_surplus"], USD),
        ("Year-1 retained loss: mean", a["mean"], USD),
        ("Year-1 retained loss: 1-in-100 (99th pct)", a["p99_1_in_100"], USD),
        ("Year-1 retained loss: 1-in-200 (99.5th pct)", a["p995_1_in_200"], USD),
        ("Year-1 retained loss: TVaR 99% (avg of worst 1%)", a["tvar99"], USD),
        ("5-year retained loss: 1-in-200", cap["five_year_retained"]["p995_1_in_200"], USD),
        ("P(surplus below regulatory minimum in any year 1-5)", cap["prob_surplus_below_min_any_year"], PCT),
        ("P(surplus negative in any year 1-5, before top-up)", cap["prob_surplus_negative_any_year"], PCT),
        ("P(parent capital call)", cap["prob_capital_call"], PCT),
        ("Capital call: 1-in-200 total", cap["capital_call_1_in_200"], USD),
        ("Premium-to-surplus, year 1", cap["premium_to_surplus_yr1"], "0.00"),
    ]
    for k, v, f in kv:
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v).number_format = f
        r += 1
    r += 1
    paths = cap["surplus_percentile_paths"]
    keys = ["0.5", "5", "25", "50", "75", "95"]
    rows = [[("inception" if y == 0 else f"end yr {y}")] + [paths[k][y] for k in keys] + [cap["min_capital"]]
            for y in range(len(paths["50"]))]
    first, last = _table(ws, r, ["Year", "0.5th pct (1-in-200)", "5th", "25th", "Median", "75th", "95th",
                                 "Regulatory minimum"], rows, {i: USD for i in range(1, 8)})
    ch = LineChart()
    ch.title = "Captive surplus percentiles (before parent top-up)"
    ch.y_axis.numFmt = '"$"#,##0'
    ch.add_data(Reference(ws, min_col=2, max_col=8, min_row=r, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ch.height, ch.width = 8, 18
    ws.add_chart(ch, f"J{r}")
    r = last + 3
    if "stress_tests" in out:
        ws.cell(row=r, column=1, value="Stress tests (losses stressed, premiums still priced on the base loss pick)").font = Font(bold=True, color=NAVY)
        rows = [[s["description"], s["annual_retained_mean"], s["annual_retained_1_in_200"],
                 s["prob_surplus_below_min_any_year"], s["prob_capital_call"], s["capital_call_1_in_200"],
                 s["saving_vs_gc_mean"]] for s in out["stress_tests"]]
        _table(ws, r + 1, ["Stress", "Mean retained (yr 1)", "1-in-200 retained (yr 1)", "P(surplus < minimum)",
                           "P(capital call)", "1-in-200 capital call", "Saving vs GC"], rows,
               {1: USD, 2: USD, 3: PCT, 4: PCT, 5: USD, 6: USD})
    _widths(ws, {"A": 52, "B": 18, "C": 18, "D": 18, "E": 16, "F": 18, "G": 16, "H": 18})

    # ----------------------------------------------------------------- Domiciles
    if "domiciles" in out:
        ws = wb.create_sheet("Domiciles")
        r = _title(ws, f"Domicile comparison for the {pid_label[focus]} structure",
                   "Regulatory figures pulled from statutes/regulators (see Basis + URL); operating-cost adjustments "
                   "and travel are assumptions. Offshore options assume an IRC 953(d) election.")
        rows = []
        for d in out["domiciles"]:
            rows.append([d["name"], d["min_capital"], d["capital_held"], d["application_fee"], d["annual_license_fee"],
                         d["premium_tax_yr1"], d["other_annual_costs"], d["pv_cost_mean"], d["saving_vs_gc_mean"],
                         d["board_meeting"]])
        first, last = _table(ws, r, ["Domicile", "Statutory min capital", "Capital held (model)", "Application fees",
                                     "Annual license fee", "Premium tax yr 1", "Other annual (assumed)",
                                     "Mean PV cost", "Saving vs GC", "Board meeting / presence"], rows,
                             {i: USD for i in range(1, 9)})
        for i in range(first, last + 1):
            ws.cell(row=i, column=10).alignment = Alignment(wrap_text=True, vertical="top")
        r = last + 2
        ws.cell(row=r, column=1, value="Citations").font = Font(bold=True, color=NAVY)
        rows = []
        for d in out["domiciles"]:
            for field_, s in d["sources"].items():
                rows.append([d["name"], field_, s["status"], s.get("value_basis", ""), s.get("url", "")])
        f2, l2 = _table(ws, r + 1, ["Domicile", "Figure", "Status", "Basis / quote", "URL"], rows)
        for i in range(f2, l2 + 1):
            ws.cell(row=i, column=3).fill = STATUS_FILL.get(ws.cell(row=i, column=3).value, STATUS_FILL["assumption"])
            ws.cell(row=i, column=4).alignment = Alignment(wrap_text=True, vertical="top")
            if ws.cell(row=i, column=5).value:
                ws.cell(row=i, column=5).hyperlink = ws.cell(row=i, column=5).value
        _widths(ws, {"A": 40, "B": 26, "C": 16, "D": 70, "E": 50, "F": 16, "G": 18, "H": 16, "I": 16, "J": 60})

    # --------------------------------------------------------------- Sensitivity
    if "tornado" in out:
        ws = wb.create_sheet("Sensitivity")
        base = out["tornado"][0]
        r = _title(ws, f"Tornado: one driver at a time, {pid_label[focus]} saving vs guaranteed cost",
                   f"Base saving vs GC {base['base_saving_vs_gc']:,.0f}; vs large deductible "
                   f"{base['base_saving_vs_ld']:,.0f}. Positive = captive cheaper.")
        rows = [[t["label"], t["path"], t["base_value"], t["low_value"], t["high_value"], t["low_saving_vs_gc"],
                 t["high_saving_vs_gc"], t["swing_vs_gc"], t["low_saving_vs_ld"], t["high_saving_vs_ld"]]
                for t in out["tornado"][1:]]
        first, last = _table(ws, r, ["Driver", "YAML path", "Base", "Low", "High", "Saving vs GC @ low",
                                     "Saving vs GC @ high", "Swing", "Saving vs LD @ low", "Saving vs LD @ high"],
                             rows, {5: USD, 6: USD, 7: USD, 8: USD, 9: USD})
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Saving vs GC at low / high value of each driver"
        ch.x_axis.numFmt = '"$"#,##0'
        ch.add_data(Reference(ws, min_col=6, max_col=7, min_row=r, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        ch.height, ch.width = 10, 20
        ws.add_chart(ch, f"A{last + 3}")
        _widths(ws, {"A": 44, "B": 44, "C": 10, "D": 10, "E": 10, "F": 18, "G": 18, "H": 14, "I": 18, "J": 18})

    # ----------------------------------------------------------------- Scenarios
    if "scenarios" in out:
        ws = wb.create_sheet("Scenarios")
        r = _title(ws, "Scenario comparison (mean PV of 5-year after-tax cost by program)",
                   "Scenarios marked 'priced on base' keep premiums at the base loss pick while true losses change.")
        ids = [p["id"] for p in progs]
        rows = []
        for s in out["scenarios"]:
            sc = cfg["scenarios"][s["id"]]
            rows.append([s["id"], s["description"] + ("" if sc.get("reprice", True) else " [priced on base]")]
                        + [s["pv_mean_by_program"][i] for i in ids]
                        + [pid_label[s["best_program"]], s["focus_saving_vs_gc"], s["focus_saving_vs_ld"],
                           s["focus_expected_loss_ratio"], str(s["focus_flags"])])
        first, last = _table(ws, r, ["Scenario", "Description"] + [pid_label[i] for i in ids]
                             + ["Lowest-cost program", f"{pid_label[focus]} saving vs GC",
                                f"{pid_label[focus]} saving vs LD", "Captive expected loss ratio", "Tax flags"],
                             rows, {j: USD for j in range(2, 2 + len(ids))}
                             | {2 + len(ids) + 1: USD, 2 + len(ids) + 2: USD, 2 + len(ids) + 3: PCT})
        for i in range(first, last + 1):
            ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        _widths(ws, {"A": 18, "B": 60} | {get_column_letter(i): 16 for i in range(3, 3 + len(ids) + 5)})

    # ----------------------------------------------------------------- Break-even
    if "breakeven" in out:
        ws = wb.create_sheet("Break-even")
        be = out["breakeven"]
        r = _title(ws, "Saving vs guaranteed cost as exposure grows (claim frequency x scale)",
                   f"{cfg['breakeven']['sims']:,} paths per point. Crossing points interpolated linearly.")
        cids = be["captive_ids"]
        rows = [[x["scale"], x["gc_premium_yr1"], x["ld_saving_vs_gc"]] + [x[f"{c}_saving_vs_gc"] for c in cids]
                for x in be["rows"]]
        first, last = _table(ws, r, ["Exposure scale", "GC premium yr 1", "Large deductible $250k"]
                             + [pid_label[c] for c in cids], rows, {i: USD for i in range(1, 3 + len(cids))})
        ch = LineChart()
        ch.title = "Mean PV saving vs guaranteed cost"
        ch.y_axis.numFmt = '"$"#,##0'
        ch.x_axis.title = "GC premium, year 1"
        ch.add_data(Reference(ws, min_col=3, max_col=3 + len(cids), min_row=r, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, min_row=first, max_row=last))
        ch.height, ch.width = 9, 20
        ws.add_chart(ch, f"A{last + 3}")
        rr = last + 23
        ws.cell(row=rr, column=1, value="Break-even (saving vs GC crosses zero)").font = Font(bold=True, color=NAVY)
        rows = []
        for c in cids:
            x = be["crossings"][c]["vs_gc"]
            y = be["crossings"][c]["vs_ld"]
            rows.append([pid_label[c], x["gc_premium_yr1"] if x else "not within range", x["scale"] if x else "",
                         y["gc_premium_yr1"] if y else "not within range"])
        _table(ws, rr + 1, ["Program", "GC premium at break-even vs GC", "Exposure scale", "Break-even vs LD"], rows,
               {1: USD, 2: "0.00", 3: USD})
        _widths(ws, {"A": 26, "B": 30, "C": 26, "D": 24, "E": 24, "F": 24, "G": 24})

    # ---------------------------------------------------------------- Sim Sample
    ws = wb.create_sheet("Sim Sample")
    r = _title(ws, "First 2,000 simulated paths", "PV of 5-year after-tax cost per program, plus the focus "
               "captive's year-1 retained loss. Use for your own pivots/histograms.")
    n = min(2000, len(next(iter(out["pv_samples"].values()))))
    ids = list(out["pv_samples"])
    rows = [[i + 1] + [float(out["pv_samples"][p][i]) for p in ids] + [float(out["retained_samples"][focus][i])]
            for i in range(n)]
    _table(ws, r, ["Path"] + [pid_label[p] for p in ids] + [f"{pid_label[focus]} yr-1 retained loss"], rows,
           {j: USD for j in range(1, len(ids) + 2)})
    ws.freeze_panes = ws.cell(row=r + 1, column=2)
    _widths(ws, {get_column_letter(i): 18 for i in range(1, len(ids) + 3)})

    # ------------------------------------------------------------------- Figures
    if figures:
        ws = wb.create_sheet("Figures")
        row = 1
        for name, p in figures.items():
            ws.cell(row=row, column=1, value=name.replace("_", " ").title()).font = Font(bold=True, color=NAVY)
            img = XLImage(str(p))
            img.width, img.height = img.width * 0.5, img.height * 0.5
            ws.add_image(img, f"A{row + 1}")
            row += int(img.height / 20) + 4

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
